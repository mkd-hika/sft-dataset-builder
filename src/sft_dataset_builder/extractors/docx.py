"""Extractor for Word (.docx), OpenDocument (.odt) and Rich Text (.rtf) files."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, ClassVar

from .base import BaseExtractor, ExtractedUnit, ExtractionResult


class DocxExtractor(BaseExtractor):
    """Handles .docx, .odt, and .rtf files."""

    supported_extensions: ClassVar[set[str]] = {".docx", ".odt", ".rtf"}

    def extract(self, path: Path, **kwargs: Any) -> ExtractionResult:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            units, warnings = self._extract_docx(path)
        elif suffix == ".odt":
            units, warnings = self._extract_odt(path)
        elif suffix == ".rtf":
            units, warnings = self._extract_rtf(path)
        else:  # pragma: no cover - guarded by can_handle
            raise ValueError(f"Unsupported extension: {suffix}")

        return ExtractionResult(
            source_path=str(path),
            source_type=suffix.lstrip("."),
            units=units,
            warnings=warnings,
        )

    # -- .docx -----------------------------------------------------------
    @staticmethod
    def _extract_docx(path: Path) -> tuple[list[ExtractedUnit], list[str]]:
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "python-docx is required for DOCX extraction. Install with `pip install python-docx`."
            ) from exc

        warnings: list[str] = []
        document = docx.Document(str(path))
        units: list[ExtractedUnit] = []

        para_index = 0
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            units.append(
                ExtractedUnit(
                    text=text,
                    metadata={"paragraph_index": para_index, "style": style},
                )
            )
            para_index += 1

        # Tables: flatten to markdown-ish rows so content isn't lost.
        for t_idx, table in enumerate(document.tables):
            rows_text = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows_text.append(" | ".join(cells))
            if rows_text:
                units.append(
                    ExtractedUnit(
                        text="\n".join(rows_text),
                        metadata={"table_index": t_idx, "kind": "table"},
                    )
                )

        return units, warnings

    # -- .odt --------------------------------------------------------------
    @staticmethod
    def _extract_odt(path: Path) -> tuple[list[ExtractedUnit], list[str]]:
        try:
            from odf import teletype  # odfpy
            from odf.opendocument import load
            from odf.text import P as ODFParagraph
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "odfpy is required for ODT extraction. Install with `pip install odfpy`."
            ) from exc

        warnings: list[str] = []
        doc = load(str(path))
        units: list[ExtractedUnit] = []
        paragraphs = doc.getElementsByType(ODFParagraph)
        for i, p in enumerate(paragraphs):
            text = teletype.extractText(p).strip()
            if text:
                units.append(ExtractedUnit(text=text, metadata={"paragraph_index": i}))
        return units, warnings

    # -- .rtf ----------------------------------------------------------------
    @staticmethod
    def _extract_rtf(path: Path) -> tuple[list[ExtractedUnit], list[str]]:
        warnings: list[str] = []
        raw = path.read_bytes().decode("utf-8", errors="replace")
        try:
            from striprtf.striprtf import rtf_to_text  # optional dependency

            plain = rtf_to_text(raw)
        except ImportError:
            warnings.append(
                "striprtf not installed; falling back to a naive RTF stripper "
                "(install `striprtf` for higher-fidelity RTF extraction)."
            )
            plain = _naive_rtf_strip(raw)

        paragraphs = [p.strip() for p in plain.split("\n") if p.strip()]
        units = [
            ExtractedUnit(text=p, metadata={"paragraph_index": i}) for i, p in enumerate(paragraphs)
        ]
        return units, warnings


def _naive_rtf_strip(raw: str) -> str:
    """Very rough RTF -> text fallback when striprtf isn't available."""
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    text = re.sub(r"{\\\*?\\[^{}]+}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\n{2,}", "\n\n", text)
    return text.strip()
