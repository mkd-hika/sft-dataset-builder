"""Extractor for PowerPoint presentations (.pptx)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, ClassVar

from .base import BaseExtractor, ExtractedUnit, ExtractionResult


class PptxExtractor(BaseExtractor):
    """Handles .pptx: one unit per slide (shape text + optional speaker notes)."""

    supported_extensions: ClassVar[set[str]] = {".pptx"}

    def extract(self, path: Path, include_notes: bool = True, **kwargs: Any) -> ExtractionResult:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "python-pptx is required for PPTX extraction. "
                "Install with `pip install python-pptx`."
            ) from exc

        presentation = Presentation(str(path))
        units: list[ExtractedUnit] = []

        for index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(text)
            if include_notes and slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = notes_frame.text.strip() if notes_frame is not None else ""
                if notes:
                    parts.append("[Notes]\n" + notes)

            slide_text = "\n\n".join(parts).strip()
            if slide_text:
                units.append(
                    ExtractedUnit(
                        text=slide_text,
                        page_number=index,
                        metadata={"slide": index},
                    )
                )

        return ExtractionResult(source_path=str(path), source_type="pptx", units=units)
